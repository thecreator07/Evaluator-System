"""Export a lesson as a Word document (.docx).

Word opens it directly. For the assignment's "document link", upload the file to
Google Drive and choose "Open with Google Docs" - that produces a shareable
Google Doc. A PDF is one click away from either: File > Download > PDF.
"""

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.shared import Pt

from models.models import Lesson


def export_lesson_docx(lesson: Lesson, output_dir: Path) -> Path:
    """Write the lesson to `output_dir` and return the file path."""

    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    path = output_dir / f"{_slug(lesson.topic)}-{stamp}.docx"
    build_lesson_document(lesson).save(path)
    return path


def build_lesson_document(lesson: Lesson) -> DocxDocument:
    document = Document()
    document.core_properties.title = lesson.title
    document.styles["Normal"].font.size = Pt(12)
    _make_schema_valid(document)

    document.add_heading(lesson.title, level=0)
    document.add_paragraph(lesson.introduction)

    _add_section(document, "What is it?", [lesson.definition])
    _add_section(document, "Why does it matter?", [lesson.why])
    _add_section(document, "How does it work?", lesson.workflow, style="List Number")
    _add_section(document, "Example", lesson.example, style="List Bullet")
    _add_key_terms(document, lesson.key_terms)
    _add_section(document, "Summary", [lesson.summary])
    return document


def _make_schema_valid(document: DocxDocument) -> None:
    """python-docx's default template omits a zoom attribute the schema requires."""
    zoom = document.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")


def _add_section(
    document: DocxDocument,
    heading: str,
    items: list[str],
    style: str | None = None,
) -> None:
    document.add_heading(heading, level=1)
    for item in items:
        document.add_paragraph(item, style=style)


def _add_key_terms(document: DocxDocument, key_terms: list[dict[str, str]]) -> None:
    document.add_heading("Key terms", level=1)
    for item in key_terms:
        term, meaning = _term_pair(item)
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(term).bold = True
        if meaning:
            paragraph.add_run(f": {meaning}")


def _term_pair(item: dict[str, str]) -> tuple[str, str]:
    """key_terms are loosely typed dicts; accept {"term", "meaning"}, {term: meaning}, or similar."""
    if "term" in item:
        return item["term"], " ".join(v for k, v in item.items() if k != "term")
    if len(item) == 1:
        ((term, meaning),) = item.items()
        return term, meaning
    values = list(item.values())
    return (values[0], " ".join(values[1:])) if values else ("", "")


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "lesson"